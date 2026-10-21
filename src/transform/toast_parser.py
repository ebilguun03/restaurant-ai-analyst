"""Turn a Toast "Product Mix" export (.xlsx) into clean, typed rows.

This file has no AWS code in it on purpose: it is pure Python, so it can be
tested on a laptop in milliseconds. The Lambda handler (handler.py) only deals
with S3 and DynamoDB and calls into this module.
"""
from __future__ import annotations

import hashlib
import re
from io import BytesIO

import openpyxl

DETAIL_SHEET = "All levels"
SUMMARY_SHEET = "Summary"

# Line types in the "Type" column of the "All levels" sheet.
#   menuItem -> a dish or drink on the menu        (revenue)
#   openItem -> a custom item rung in by staff     (revenue)
#   giftCard -> a gift card sold                   (NOT revenue: Toast leaves it
#               out of its own total, because the money is owed back as food later)
REVENUE_TYPES = {"menuItem", "openItem"}
KEEP_TYPES = REVENUE_TYPES | {"giftCard"}

REQUIRED_COLUMNS = [
    "Type", "Menu", "Menu group", "Subgroup", "Item, open item",
    "Qty sold", "Avg. price", "Gross sales", "Discount amount",
    "Refund amount", "Void amount", "Net sales", "Tax",
]

TITLE_RE = re.compile(r"ProductMix_(\d{4}-\d{2})-\d{2}_(\d{4}-\d{2})-\d{2}")

# How far our sum may drift from Toast's own total before we refuse the file.
# Toast rounds to cents, so a few cents of drift is normal.
RECONCILE_TOLERANCE = 0.05


class ParseError(ValueError):
    """The file is not a Product Mix export we understand, or its numbers don't add up."""


def _num(value) -> float:
    if value in (None, ""):
        return 0.0
    return float(value)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def make_item_key(menu: str, group: str, item: str) -> str:
    """A stable ID for a menu item, so the same dish matches across months."""
    raw = f"{menu}|{group}|{item}".lower()
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]


def month_from_title(title: str) -> str:
    match = TITLE_RE.search(title)
    if not match:
        raise ParseError(f"Can't find a ProductMix date range in {title!r}")
    start_month, end_month = match.groups()
    if start_month != end_month:
        raise ParseError(
            f"Export covers {start_month} to {end_month}; export one calendar month per file"
        )
    return start_month


def parse_location(line: str) -> tuple[str, str]:
    """'-2120 S Highland Dr, Suite #15 -Ozora Izakaya Japanese Bar' -> (address, name)."""
    cleaned = line.strip().lstrip("-").strip()
    if " -" in cleaned:
        address, name = cleaned.rsplit(" -", 1)
        return address.strip(), name.strip()
    return "", cleaned


def parse_product_mix(data: bytes, source_name: str) -> dict:
    """Parse one export. Returns {"restaurant": {...}, "month": "YYYY-MM", "rows": [...], "check": {...}}."""
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # corrupt file, wrong format, etc.
        raise ParseError(f"{source_name} is not a readable .xlsx file: {exc}") from exc

    for sheet in (SUMMARY_SHEET, DETAIL_SHEET):
        if sheet not in workbook.sheetnames:
            raise ParseError(f"{source_name} has no {sheet!r} sheet; is it a Toast Product Mix export?")

    summary = list(workbook[SUMMARY_SHEET].iter_rows(values_only=True, max_row=2))
    title = _text(summary[0][0]) if summary else ""
    try:
        month = month_from_title(title)
    except ParseError:
        month = month_from_title(source_name)  # fall back to the file name
    address, name = parse_location(_text(summary[1][0]) if len(summary) > 1 else "")
    restaurant_id = slugify(name) or "unknown-restaurant"

    rows_iter = workbook[DETAIL_SHEET].iter_rows(values_only=True)
    header = [_text(h) for h in next(rows_iter)]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ParseError(f"{source_name} is missing columns: {missing}")
    col = {name_: i for i, name_ in enumerate(header)}

    toast_total = None
    rows = []
    for raw in rows_iter:
        line_type = _text(raw[col["Type"]])
        if line_type == "" and toast_total is None and _text(raw[col["Menu"]]) == "":
            # First blank-type row with no menu is Toast's grand total. Keep it to check against.
            toast_total = {"qty": _num(raw[col["Qty sold"]]), "net": _num(raw[col["Net sales"]])}
            continue
        if line_type not in KEEP_TYPES:
            continue  # subtotal rows for menus and groups; we recompute those ourselves
        menu = _text(raw[col["Menu"]])
        group = _text(raw[col["Menu group"]])
        item = _text(raw[col["Item, open item"]])
        if not item:
            continue
        rows.append({
            "restaurant_id": restaurant_id,
            "month": month,
            "item_key": make_item_key(menu, group, item),
            "item_name": item,
            "menu": menu,
            "menu_group": group,
            "subgroup": _text(raw[col["Subgroup"]]),
            "line_type": line_type,
            "is_revenue": line_type in REVENUE_TYPES,
            "qty_sold": int(round(_num(raw[col["Qty sold"]]))),
            "avg_price": round(_num(raw[col["Avg. price"]]), 2),
            "gross_sales": round(_num(raw[col["Gross sales"]]), 2),
            "discount_amount": round(_num(raw[col["Discount amount"]]), 2),
            "refund_amount": round(_num(raw[col["Refund amount"]]), 2),
            "void_amount": round(_num(raw[col["Void amount"]]), 2),
            "net_sales": round(_num(raw[col["Net sales"]]), 2),
            "tax": round(_num(raw[col["Tax"]]), 2),
            "source_file": source_name,
        })

    if not rows:
        raise ParseError(f"{source_name} has no item rows")

    check = reconcile(rows, toast_total)
    return {
        "restaurant": {"restaurant_id": restaurant_id, "name": name, "address": address},
        "month": month,
        "rows": rows,
        "check": check,
    }


def reconcile(rows: list[dict], toast_total: dict | None) -> dict:
    """Data-quality gate: our revenue rows must add up to Toast's own grand total."""
    our_qty = sum(r["qty_sold"] for r in rows if r["is_revenue"])
    our_net = round(sum(r["net_sales"] for r in rows if r["is_revenue"]), 2)
    if toast_total is None:
        raise ParseError("No grand-total row found, so the file can't be checked")
    net_diff = round(abs(our_net - toast_total["net"]), 2)
    qty_diff = abs(our_qty - toast_total["qty"])
    if net_diff > RECONCILE_TOLERANCE or qty_diff > 0:
        raise ParseError(
            f"Totals don't match Toast: net {our_net} vs {toast_total['net']}, "
            f"qty {our_qty} vs {toast_total['qty']}"
        )
    return {"rows": len(rows), "qty_sold": our_qty, "net_sales": our_net, "net_diff": net_diff}
